import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook

wb = load_workbook(filename='result.xlsx')
ws = wb.active
tests = ["test" + str(i) for i in range(1, 6)]
result = {}

for c in range(2, 7):
    algo = ws.cell(row=1, column=c).value
    result[algo] = []
    for r in range(2, 7):
        result[algo].append(ws.cell(row=r, column=c).value)

x = np.arange(len(tests))  # the label locations
width = 0.19  # the width of the bars
multiplier = 0

plt.rcParams.update({'font.size': 16})
fig, ax = plt.subplots(layout='constrained', figsize=(18, 9))

for algo, exec_time in result.items():
    offset = width * multiplier
    rects = ax.bar(x + offset, exec_time, width, label=algo)
    ax.bar_label(rects, padding=3)
    multiplier += 1

# Add some text for labels, title and custom x-axis tick labels, etc.
ax.set_ylabel('Thời gian thực hiện (ms)')
ax.set_title('Kết quả thử nghiệm bộ dữ liệu')
ax.set_xticks(x + width*2, tests)
ax.legend(loc='right', bbox_to_anchor=(1.05, 0.5), ncol=1)
ax.set_ylim(0, 8200)

plt.show()

